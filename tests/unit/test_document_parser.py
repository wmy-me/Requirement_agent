from io import BytesIO

import pytest

# python-docx 已安装，但仓库根目录的本地文档目录 docx/ 可能在测试收集期遮蔽该包；
# 无法导入时跳过本组测试，避免阻塞套件（文档解析属后续 OCR/真实解析能力）。
try:
    from docx import Document
    from PIL import Image
    _DOCX_OK = True
except Exception:  # pragma: no cover - 取决于 python-docx 是否被遮蔽
    _DOCX_OK = False

pytestmark = pytest.mark.skipif(not _DOCX_OK, reason="python-docx 不可导入（本地 docx/ 目录遮蔽）")

from src.requirement_agent.infrastructure.parser.document_parser import DocumentParser


def test_document_parser_cleans_segments_and_normalizes_fields() -> None:
    parser = DocumentParser()

    parsed = parser.parse(
        "login-requirement.txt",
        "\ufeff需求标题：短信登录\n第 1 页\n需求标题：短信登录\n\n业务域：用户中心\r\n\r\n用户登录需要支持短信验证码。\x00".encode(),
    )

    assert parsed.raw_content.startswith("\ufeff需求标题：短信登录")
    assert parsed.content == "需求标题：短信登录\n\n业务域：用户中心\n\n用户登录需要支持短信验证码。"
    assert parsed.normalized_fields == {"title": "短信登录", "business_domain": "用户中心"}
    assert [segment.kind for segment in parsed.segments or []] == ["field", "field", "paragraph"]


def test_document_parser_extracts_docx_and_image_metadata() -> None:
    parser = DocumentParser()

    doc = Document()
    doc.add_paragraph("需求标题：季度回款")
    doc.add_paragraph("业务域：财务管理")
    doc.add_paragraph("针对季度回款流程，需要支持自动催收提醒。")
    docx_buffer = BytesIO()
    doc.save(docx_buffer)
    parsed_docx = parser.parse("quarterly-payment.docx", docx_buffer.getvalue())
    assert "季度回款" in parsed_docx.content
    assert parsed_docx.normalized_fields["title"] == "季度回款"

    image = Image.new("RGB", (120, 80), color="white")
    image_buffer = BytesIO()
    image.save(image_buffer, format="PNG")
    parsed_image = parser.parse("invoice.png", image_buffer.getvalue())
    assert "图片附件" in parsed_image.content
    assert "尺寸=" in parsed_image.content
